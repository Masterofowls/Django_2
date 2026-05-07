import os
import shutil
from datetime import date, datetime
import openpyxl
from django.conf import settings
from django.core.management.base import BaseCommand
from users.models import User
from mainapp.models import (
    Category,
    Manufacturer,
    Order,
    PickUpPoint,
    Product,
    Supplier,
)

def _parse_date(value):
    if value is None:
        return None
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


class Command(BaseCommand):
    help = "импорт данных xlsx"

    def handle(self, *args, **options):
        self.import_dir = os.path.join(settings.BASE_DIR, "import")
        self._import_users()
        self._import_pickups()
        self._import_orders()
        self._import_products()

    def _import_users(self):
        file_path = os.path.join(self.import_dir, "user_import.xlsx")
        if not os.path.isfile(file_path):
            self.stderr.write(self.style.WARNING(f"Файл не найден: {file_path}"))
            return
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        role_mapping = {
            "администратор": "admin",
            "менеджер": "manager",
            "авторизированный клиент": "client",
        }
        created = 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            role_raw = row[0]
            if not role_raw:
                continue
            role_raw = str(role_raw).strip().lower()
            fullname = str(row[1] or "").strip()
            username = str(row[2] or "").strip()
            password = str(row[3] or "").strip()
            role = role_mapping.get(role_raw, "client")
            user, is_new = User.objects.get_or_create(
                username=username, defaults={"full_name": fullname, "role": role}
            )
            if is_new:
                user.set_password(password)
                user.save()
                created += 1
        self.stdout.write(self.style.SUCCESS(f"Создано пользователей: {created}"))


    def _import_pickups(self):
        file_path = os.path.join(self.import_dir, "Пункты выдачи_import.xlsx")
        if not os.path.isfile(file_path):
            self.stderr.write(self.style.WARNING(f"Файл не найден: {file_path}"))
            return
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        created = 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            address = row[0]
            if not address:
                continue
            address = str(address).strip()
            _, is_new = PickUpPoint.objects.get_or_create(address=address)
            if is_new:
                created += 1
        self.stdout.write(self.style.SUCCESS(f"Создано пунктов выдачи: {created}"))

    def _import_orders(self):
        file_path = os.path.join(self.import_dir, "Заказ_import.xlsx")
        if not os.path.isfile(file_path):
            self.stderr.write(self.style.WARNING(f"Файл не найден: {file_path}"))
            return
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        status_mapping = {
            "новый": "new",
            "завершён": "compleated",
            "завершен": "compleated",
            "отменённый": "canceled",
            "отмененный": "canceled",
        }
        created = 0
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            number = row[0]
            if not number:
                continue
            try:
                number = int(number)
            except (ValueError, TypeError):
                self.stderr.write(
                    self.style.WARNING(
                        f"Строка {idx}: неверный номер заказа '{number}'"
                    )
                )
                continue
            order_date = _parse_date(row[2])
            if order_date is None:
                self.stderr.write(
                    self.style.WARNING(
                        f"Строка {idx}: ошибка даты  заказа '{row[2]}'"
                    )
                )
                continue
            delivery_date = _parse_date(row[3])
            pickup_address = str(row[4] or "").strip()
            status_raw = str(row[7] or "новый").strip().lower()
            status = status_mapping.get(status_raw, "new")

            pickup_point = None
            if pickup_address:
                pickup_point, _ = PickUpPoint.objects.get_or_create(
                    address=pickup_address
                )

            _, is_new = Order.objects.get_or_create(
                number=number,
                defaults={
                    "order_date": order_date,
                    "delivery_date": delivery_date,
                    "pickup_point": pickup_point,
                    "status": status,
                },
            )
            if is_new:
                created += 1
        self.stdout.write(self.style.SUCCESS(f"Создано заказов: {created}"))

    def _import_products(self):
        file_path = os.path.join(self.import_dir, "Tovar.xlsx")
        if not os.path.isfile(file_path):
            self.stderr.write(self.style.WARNING(f"Файл не найден: {file_path}"))
            return
        wb = openpyxl.load_workbook(file_path)
        ws = wb.active
        created = 0
        photo_path = ""
        products_media_dir = os.path.join(settings.MEDIA_ROOT, 'products')
        os.makedirs(products_media_dir, exist_ok = True)
        for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
            self.stdout.write(f'{row}')
            article = row[0]
            if not article:
                continue
            article = str(article).strip()
            name = str(row[1] or "").strip()
            description = str(row[2] or "").strip()
            try:
                price = float(row[3] or 0)
            except (ValueError, TypeError):
                self.stderr.write(
                    self.style.WARNING(f"Строка {idx}: некорректная цена")
                )
                continue
            unit = str(row[4] or "шт.").strip()
            try:
                stock_quantity = int(row[8] or 0)
                self.stdout.write(f'{row[8]}')
            except (ValueError, TypeError):
             stock_quantity = 0
            try:
                discount = float(row[7] or 0)
            except (ValueError, TypeError):
                discount = 0
            category_name = str(row[6] or "").strip()
            manufacturer_name = str(row[5] or "").strip()
            supplier_name = str(row[4] or "").strip()
            photo_name = str(row[10] or "").strip()
            if photo_name :
              src_photo = os.path.join(self.import_dir, photo_name)
              if os.path.isfile(src_photo):
                dst_photo = os.path.join(products_media_dir, photo_name)
                shutil.copy2(src_photo, dst_photo)
                photo_path = f'products/{photo_name}'



            self.stdout.write(
                f"{article}, {name}, {description}, {price}, {unit}, {stock_quantity}, {discount}, {category_name}, {manufacturer_name}, {supplier_name}"
            )

            if (
                not name
                or not category_name
                or not manufacturer_name
                or not supplier_name
            ):
                self.stderr.write(
                    self.style.WARNING(f"Строка {idx}: пропущены обязательные поля")
                )
                continue

            category, _ = Category.objects.get_or_create(name=category_name)
            manufacturer, _ = Manufacturer.objects.get_or_create(name=manufacturer_name)
            supplier, _ = Supplier.objects.get_or_create(name=supplier_name)

            _, is_new = Product.objects.get_or_create(
                article=article,
                defaults={
                    "name": name,
                    "description": description,
                    "price": price,
                    "unit": unit,
                    "stock_quantity": stock_quantity,
                    "discount": discount,
                    "category": category,
                    "manufacturer": manufacturer,
                    "supplier": supplier,
                   "photo": photo_path if photo_path else None,
                },
            )
            if is_new:
                created += 1
        self.stdout.write(self.style.SUCCESS(f"Создано товаров: {created}"))
